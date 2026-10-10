"""Extract plain text from office document bytes for Gemini summarization.

Only stdlib + pypdf / python-docx / python-pptx / openpyxl. Raises
HTTPException(400) with a user-facing message when the file has no
extractable text (e.g. scanned PDFs) or is corrupt.
"""

import io

from fastapi import HTTPException


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception:
        raise HTTPException(
            status_code=400, detail="File is not a valid PDF")
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise HTTPException(
                status_code=400,
                detail="Encrypted PDFs are not supported",
            )
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _extract_docx(content: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
    except Exception:
        raise HTTPException(
            status_code=400, detail="File is not a valid DOCX document")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _extract_pptx(content: bytes) -> str:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(content))
    except Exception:
        raise HTTPException(
            status_code=400, detail="File is not a valid PPTX presentation")
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
    return "\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(
            status_code=400, detail="File is not a valid XLSX spreadsheet")
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
}


def extract_document_text(content: bytes, ext: str) -> str:
    """Extract text from an office document. Raises HTTPException on failure."""
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        raise HTTPException(
            status_code=400, detail=f"Unsupported document type '{ext}'")
    text = extractor(content)
    if not text or not text.strip():
        if ext == ".pdf":
            raise HTTPException(
                status_code=400,
                detail="PDF contains no extractable text "
                "(scanned images without OCR are not supported)",
            )
        raise HTTPException(
            status_code=400, detail="Document contains no readable text")
    return text
